"""Eksport obsady do Excela i PDF oraz porównanie z opublikowaną wersją."""

from __future__ import annotations

import datetime as dt
import io
from collections import defaultdict
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Assignment, Match, MatchStatus, RefereeRole

ROLE_LABELS: dict[str, str] = {
    RefereeRole.MAIN.value: "Sędzia główny",
    RefereeRole.ASSISTANT_1.value: "Asystent 1",
    RefereeRole.ASSISTANT_2.value: "Asystent 2",
    RefereeRole.FOURTH.value: "Sędzia techniczny",
    RefereeRole.OBSERVER.value: "Obserwator",
}
ROLE_SHORT: dict[str, str] = {
    RefereeRole.MAIN.value: "G",
    RefereeRole.ASSISTANT_1.value: "A1",
    RefereeRole.ASSISTANT_2.value: "A2",
    RefereeRole.FOURTH.value: "T",
    RefereeRole.OBSERVER.value: "O",
}
MISSING = "— brak —"
WEEKDAYS = ("poniedziałek", "wtorek", "środa", "czwartek", "piątek", "sobota", "niedziela")
FONT_DIR = Path(__file__).resolve().parents[1] / "assets" / "fonts"


def _venue_address(m: Match) -> str:
    v = m.venue
    city = " ".join(p for p in (v.postal_code, v.city) if p)
    return ", ".join(p for p in (v.address, city) if p)


def collect(db: Session, date_from: dt.date, date_to: dt.date) -> list[dict[str, Any]]:
    """Wiersze obsady okresu — wspólne źródło dla Excela, PDF i publikacji (JSON-owalne)."""
    matches = db.scalars(
        select(Match)
        .options(selectinload(Match.venue), selectinload(Match.assignments).selectinload(Assignment.referee))
        .where(Match.date >= date_from, Match.date <= date_to)
        .order_by(Match.date, Match.start_time, Match.id)
    )
    rows = []
    for m in matches:
        by_role = {a.role: a for a in m.assignments}
        crew = []
        for role in m.required_roles + [r for r in by_role if r not in m.required_roles]:
            a = by_role.get(role)
            crew.append(
                {
                    "role": role.value,
                    "referee_id": a.referee_id if a else None,
                    "referee_name": a.referee.full_name if a and a.referee else None,
                    "distance_km": a.distance_km if a else None,
                }
            )
        rows.append(
            {
                "match_id": m.id,
                "date": m.date.isoformat(),
                "start_time": m.start_time.strftime("%H:%M"),
                "league": m.league,
                "home_team": m.home_team,
                "away_team": m.away_team,
                "venue_name": m.venue.name,
                "venue_address": _venue_address(m),
                "cancelled": m.status == MatchStatus.CANCELLED,
                "crew": crew,
            }
        )
    return rows


def _signature(row: dict[str, Any]) -> tuple:
    """To, co widzi sędzia: termin, miejsce, skład. Odległości nie liczą się jako zmiana."""
    return (
        row["date"],
        row["start_time"],
        row["venue_name"],
        row["cancelled"],
        tuple((c["role"], c["referee_id"]) for c in row["crew"]),
    )


def changed_matches(published: list[dict[str, Any]], current: list[dict[str, Any]]) -> list[int]:
    """Mecze, które od publikacji zmieniły skład/termin, doszły lub zniknęły."""
    old = {r["match_id"]: _signature(r) for r in published}
    new = {r["match_id"]: _signature(r) for r in current}
    return sorted(mid for mid in old.keys() | new.keys() if old.get(mid) != new.get(mid))


def _day_label(iso: str) -> str:
    d = dt.date.fromisoformat(iso)
    return f"{WEEKDAYS[d.weekday()]} {d.strftime('%d.%m.%Y')}"


def _crew_name(c: dict[str, Any]) -> str:
    return c["referee_name"] or MISSING


def filename(prefix: str, date_from: dt.date, date_to: dt.date, ext: str) -> str:
    span = date_from.isoformat() if date_from == date_to else f"{date_from.isoformat()}_{date_to.isoformat()}"
    return f"{prefix}_{span}.{ext}"


# --- Excel ------------------------------------------------------------------------
def to_xlsx(rows: list[dict[str, Any]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Obsada"
    roles = [r.value for r in RefereeRole if any(c["role"] == r.value for row in rows for c in row["crew"])]
    if not roles:
        roles = [RefereeRole.MAIN.value]
    header = ["Data", "Dzień", "Godzina", "Liga", "Gospodarz", "Gość", "Obiekt", "Adres"]
    header += [ROLE_LABELS[r] for r in roles] + ["Uwagi"]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(header)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    missing_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in ws[2]:
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        names = {c["role"]: c["referee_name"] for c in row["crew"]}
        required = {c["role"] for c in row["crew"]}
        d = dt.date.fromisoformat(row["date"])
        values = [d, WEEKDAYS[d.weekday()], row["start_time"], row["league"], row["home_team"], row["away_team"],
                  row["venue_name"], row["venue_address"]]
        values += [(names.get(r) or MISSING) if r in required else "" for r in roles]
        values.append("ODWOŁANY" if row["cancelled"] else "")
        ws.append(values)
        excel_row = ws.max_row
        ws.cell(excel_row, 1).number_format = "DD.MM.YYYY"
        for i, r in enumerate(roles):
            if r in required and not names.get(r):
                ws.cell(excel_row, 9 + i).fill = missing_fill
    widths = [11, 13, 8, 22, 28, 28, 30, 38] + [24] * len(roles) + [12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(header))}{max(ws.max_row, 2)}"

    # Arkusz z rozpiską na sędziów — do rozesłania i do kilometrówki.
    per_ref: dict[str, list[tuple[dict, dict]]] = defaultdict(list)
    for row in rows:
        for c in row["crew"]:
            if c["referee_name"]:
                per_ref[c["referee_name"]].append((row, c))
    ws2 = wb.create_sheet("Sędziowie")
    ws2.append(["Sędzia", "Data", "Godzina", "Rola", "Mecz", "Obiekt", "Dojazd (km)"])
    for cell in ws2[1]:
        cell.font = bold
        cell.fill = fill
    for name in sorted(per_ref):
        total = 0.0
        for row, c in per_ref[name]:
            km = c["distance_km"]
            total += km or 0
            ws2.append([name, dt.date.fromisoformat(row["date"]), row["start_time"], ROLE_LABELS[c["role"]],
                        f"{row['home_team']} – {row['away_team']}", row["venue_name"],
                        round(km, 1) if km is not None else None])
            ws2.cell(ws2.max_row, 2).number_format = "DD.MM.YYYY"
        ws2.append([f"{name} — razem", None, None, None, f"{len(per_ref[name])} mecz(e)", None, round(total, 1)])
        for cell in ws2[ws2.max_row]:
            cell.font = Font(bold=True)
    for i, w in enumerate([26, 11, 8, 18, 46, 30, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------------
_fonts_registered = False


def _register_fonts() -> None:
    global _fonts_registered
    if _fonts_registered:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    _fonts_registered = True


def to_pdf(rows: list[dict[str, Any]], title: str, subtitle: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    base = ParagraphStyle("base", fontName="DejaVu", fontSize=8, leading=10)
    bold = ParagraphStyle("bold", parent=base, fontName="DejaVu-Bold")
    missing = ParagraphStyle("missing", parent=base, textColor=colors.HexColor("#B42318"))
    h1 = ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=15, leading=19)
    h2 = ParagraphStyle("h2", fontName="DejaVu-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    small = ParagraphStyle("small", parent=base, textColor=colors.HexColor("#555555"))

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list = [Paragraph(esc(title), h1)]
    if subtitle:
        story.append(Paragraph(esc(subtitle), small))
    story.append(Spacer(1, 4 * mm))
    if not rows:
        story.append(Paragraph("Brak meczów w wybranym okresie.", base))

    by_day: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_day[row["date"]].append(row)
    widths = [14 * mm, 30 * mm, 70 * mm, 62 * mm, 101 * mm]
    for day in sorted(by_day):
        story.append(Paragraph(_day_label(day).capitalize(), h2))
        data = [[Paragraph(h, bold) for h in ("Godz.", "Liga", "Mecz", "Obiekt", "Obsada")]]
        for row in by_day[day]:
            crew_lines = []
            for c in row["crew"]:
                label = ROLE_SHORT.get(c["role"], c["role"])
                name = esc(c["referee_name"]) if c["referee_name"] else f'<font color="#B42318">{MISSING}</font>'
                crew_lines.append(f"<b>{label}</b>: {name}")
            match = f"{esc(row['home_team'])} – {esc(row['away_team'])}"
            if row["cancelled"]:
                match += ' <font color="#B42318">(ODWOŁANY)</font>'
            venue = esc(row["venue_name"])
            if row["venue_address"]:
                venue += f'<br/><font color="#555555">{esc(row["venue_address"])}</font>'
            data.append(
                [
                    Paragraph(row["start_time"], base),
                    Paragraph(esc(row["league"]), base),
                    Paragraph(match, base),
                    Paragraph(venue, base),
                    Paragraph("<br/>".join(crew_lines) or MISSING, base if crew_lines else missing),
                ]
            )
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4D0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FB")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVu", 7)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(12 * mm, 7 * mm, f"Obsada+ · {title}")
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 7 * mm, f"Strona {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
        bottomMargin=14 * mm, title=title, author="Obsada+",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def period_title(date_from: dt.date, date_to: dt.date) -> str:
    if date_from == date_to:
        return f"Obsada sędziowska — {_day_label(date_from.isoformat())}"
    return f"Obsada sędziowska — {date_from.strftime('%d.%m')}–{date_to.strftime('%d.%m.%Y')}"
