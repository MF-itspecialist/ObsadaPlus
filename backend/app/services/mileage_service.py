"""Statystyki okresu (miesiąc, sezon) i zestawienie kilometrówki sędziów.

Kilometry liczone są dla każdego dnia osobno: dom → mecz 1 → mecz 2 → … → dom.
Powrót do domu przyjmujemy jako tę samą trasę co dojazd z domu na ostatni obiekt
(te trasy są już w pamięci tras po generowaniu, więc nie ma nowych zapytań do map).
"""

from __future__ import annotations

import datetime as dt
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Assignment, Match, MatchStatus, RefereeRole
from app.optimization.licenses import level_name
from app.services.routing_service import GeoPoint, RoutingService
from app.services.settings_service import get_settings, routing_for

MAIN_ROLES = {RefereeRole.MAIN}
ASSISTANT_ROLES = {RefereeRole.ASSISTANT_1, RefereeRole.ASSISTANT_2}


@dataclass
class RefereeTotals:
    referee_id: int
    name: str
    location: str
    matches: int = 0
    as_main: int = 0
    as_assistant: int = 0
    days: set[dt.date] = field(default_factory=set)
    km: float = 0.0
    levels: Counter = field(default_factory=Counter)
    trips: list[dict[str, Any]] = field(default_factory=list)


def season_bounds(day: dt.date) -> tuple[dt.date, dt.date]:
    """Sezon piłkarski: 1 lipca – 30 czerwca."""
    start_year = day.year if day.month >= 7 else day.year - 1
    return dt.date(start_year, 7, 1), dt.date(start_year + 1, 6, 30)


def _assignments(db: Session, date_from: dt.date, date_to: dt.date) -> list[Assignment]:
    return list(
        db.scalars(
            select(Assignment)
            .join(Match)
            .options(selectinload(Assignment.referee), selectinload(Assignment.match).selectinload(Match.venue))
            .where(Match.date >= date_from, Match.date <= date_to, Match.status != MatchStatus.CANCELLED)
            .order_by(Match.date, Match.start_time)
        )
    )


def _day_route(routing: RoutingService, home: GeoPoint, venues: list[GeoPoint]) -> list[float]:
    """Kilometry odcinków dnia: dom → v1 → … → vn → dom (powrót = dojazd na vn z domu)."""
    legs, here = [], home
    for v in venues:
        legs.append(routing.get_distance(here, v))
        here = v
    legs.append(routing.get_distance(home, here) if venues else 0.0)
    return legs


def compute(db: Session, date_from: dt.date, date_to: dt.date, routing: RoutingService | None = None) -> dict[str, Any]:
    settings = get_settings(db)
    routing = routing or routing_for(settings, db)
    rate = settings.mileage_rate
    assignments = _assignments(db, date_from, date_to)

    totals: dict[int, RefereeTotals] = {}
    by_day: dict[tuple[int, dt.date], list[Assignment]] = defaultdict(list)
    for a in assignments:
        r = a.referee
        t = totals.get(r.id) or totals.setdefault(r.id, RefereeTotals(r.id, r.full_name, r.location_label))
        t.matches += 1
        t.as_main += a.role in MAIN_ROLES
        t.as_assistant += a.role in ASSISTANT_ROLES
        t.days.add(a.match.date)
        t.levels[level_name(a.match.league) or (a.match.league or "inne")] += 1
        by_day[(r.id, a.match.date)].append(a)

    for (rid, day), items in sorted(by_day.items(), key=lambda kv: kv[0][1]):
        ref = items[0].referee
        home = GeoPoint(ref.latitude, ref.longitude, ref.address or ref.city)
        # Ten sam mecz w dwóch rolach się nie zdarza, ale ten sam obiekt dwa razy z rzędu = 0 km.
        venues = [GeoPoint(a.match.venue.latitude, a.match.venue.longitude, a.match.venue.name) for a in items]
        legs = _day_route(routing, home, venues)
        km = round(sum(legs), 1)
        totals[rid].km += km
        totals[rid].trips.append(
            {
                "date": day,
                "matches": [f"{a.match.start_time:%H:%M} {a.match.home_team} – {a.match.away_team}" for a in items],
                "route": " → ".join(["dom", *[a.match.venue.name for a in items], "dom"]),
                "km": km,
                "amount": round(km * rate, 2),
            }
        )

    referees = sorted(totals.values(), key=lambda t: (-t.matches, t.name))
    matches = list(
        db.scalars(
            select(Match)
            .options(selectinload(Match.assignments))
            .where(Match.date >= date_from, Match.date <= date_to, Match.status != MatchStatus.CANCELLED)
        )
    )
    slots_total = sum(len(m.required_roles) for m in matches)
    slots_filled = sum(len({a.role for a in m.assignments} & set(m.required_roles)) for m in matches)
    by_month: dict[str, dict[str, float]] = defaultdict(lambda: {"matches": 0, "slots_filled": 0, "km": 0.0})
    for m in matches:
        key = f"{m.date:%Y-%m}"
        by_month[key]["matches"] += 1
        by_month[key]["slots_filled"] += len({a.role for a in m.assignments} & set(m.required_roles))
    for t in referees:
        for trip in t.trips:
            by_month[f"{trip['date']:%Y-%m}"]["km"] += trip["km"]
    levels = Counter(level_name(m.league) or (m.league or "inne") for m in matches)
    km_total = round(sum(t.km for t in referees), 1)

    return {
        "date_from": date_from,
        "date_to": date_to,
        "mileage_rate": rate,
        "summary": {
            "matches": len(matches),
            "matches_fully_assigned": sum(
                1 for m in matches if set(m.required_roles) <= {a.role for a in m.assignments}
            ),
            "slots_total": slots_total,
            "slots_filled": slots_filled,
            "referees_used": len(referees),
            "km_total": km_total,
            "amount_total": round(km_total * rate, 2),
        },
        "by_month": [{"month": k, **{kk: round(vv, 1) for kk, vv in v.items()}} for k, v in sorted(by_month.items())],
        "by_level": [{"level": k, "matches": v} for k, v in levels.most_common()],
        "referees": [
            {
                "referee_id": t.referee_id,
                "name": t.name,
                "location": t.location,
                "matches": t.matches,
                "as_main": t.as_main,
                "as_assistant": t.as_assistant,
                "days": len(t.days),
                "km": round(t.km, 1),
                "amount": round(t.km * rate, 2),
                "levels": dict(t.levels.most_common()),
                "trips": t.trips,
            }
            for t in referees
        ],
    }


def to_xlsx(data: dict[str, Any]) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    rate = data["mileage_rate"]
    period = f"{data['date_from']:%d.%m.%Y} – {data['date_to']:%d.%m.%Y}"

    def header(ws, cols, widths):
        ws.append(cols)
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = head_font, head_fill
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{ws.max_row + 1}"  # ws.cell() utworzyłby pusty wiersz

    wb = Workbook()
    ws = wb.active
    ws.title = "Kilometrówka"
    ws.append([f"Kilometrówka sędziów {period} (stawka {rate:.4f} zł/km)"])
    ws["A1"].font = Font(bold=True, size=13)
    header(ws, ["Sędzia", "Miejscowość", "Mecze", "Jako główny", "Jako asystent", "Dni", "Km", "Kwota (zł)"],
           [28, 24, 8, 12, 13, 6, 10, 12])
    for r in data["referees"]:
        ws.append([r["name"], r["location"], r["matches"], r["as_main"], r["as_assistant"], r["days"], r["km"], r["amount"]])
        ws.cell(ws.max_row, 8).number_format = "#,##0.00"
    s = data["summary"]
    ws.append(["Razem", None, sum(r["matches"] for r in data["referees"]), None, None, None, s["km_total"], s["amount_total"]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.cell(ws.max_row, 8).number_format = "#,##0.00"

    ws2 = wb.create_sheet("Przejazdy")
    header(ws2, ["Sędzia", "Data", "Trasa", "Mecze", "Km", "Kwota (zł)"], [28, 11, 60, 60, 8, 11])
    for r in data["referees"]:
        for t in r["trips"]:
            ws2.append([r["name"], t["date"], t["route"], "; ".join(t["matches"]), t["km"], t["amount"]])
            ws2.cell(ws2.max_row, 2).number_format = "DD.MM.YYYY"
            ws2.cell(ws2.max_row, 6).number_format = "#,##0.00"
    ws2.auto_filter.ref = f"A1:F{max(ws2.max_row, 1)}"

    ws3 = wb.create_sheet("Miesiące")
    header(ws3, ["Miesiąc", "Mecze", "Obsadzone miejsca", "Km"], [10, 8, 18, 10])
    for m in data["by_month"]:
        ws3.append([m["month"], m["matches"], m["slots_filled"], m["km"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
