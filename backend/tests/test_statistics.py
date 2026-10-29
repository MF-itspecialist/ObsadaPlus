"""Statystyki okresu i kilometrówka."""

import datetime as dt
import io

from openpyxl import load_workbook

from app.services.mileage_service import _day_route, season_bounds
from tests.factories import LEGIONOWO, PIASECZNO, WARSZAWA, FixedTimeRouting

PERIOD = {"date_from": "2026-10-09", "date_to": "2026-10-12"}


def test_day_route_goes_home_after_last_match():
    routing = FixedTimeRouting(minutes=20, km=10)
    assert _day_route(routing, WARSZAWA, [LEGIONOWO]) == [10, 10]
    assert _day_route(routing, WARSZAWA, [LEGIONOWO, PIASECZNO]) == [10, 10, 10]
    # dwa mecze na tym samym obiekcie: między nimi 0 km
    assert _day_route(routing, WARSZAWA, [LEGIONOWO, LEGIONOWO]) == [10, 0, 10]
    assert _day_route(routing, WARSZAWA, []) == [0.0]


def test_season_runs_from_july_to_june():
    assert season_bounds(dt.date(2026, 10, 10)) == (dt.date(2026, 7, 1), dt.date(2027, 6, 30))
    assert season_bounds(dt.date(2027, 3, 1)) == (dt.date(2026, 7, 1), dt.date(2027, 6, 30))


def test_statistics_and_mileage_export(seeded_client):
    c = seeded_client
    preview = c.post("/api/optimization/generate", json=PERIOD).json()
    c.post("/api/optimization/confirm", json={"run_id": preview["run_id"]})
    c.put("/api/settings", json={**c.get("/api/settings").json(), "mileage_rate": 1.0})

    data = c.get("/api/statistics", params=PERIOD).json()
    s = data["summary"]
    assert s["matches"] == 24
    assert s["slots_filled"] == preview["stats"]["slots_assigned"]
    assert s["referees_used"] == len(data["referees"]) > 0
    assert sum(r["matches"] for r in data["referees"]) == s["slots_filled"]
    assert abs(sum(r["km"] for r in data["referees"]) - s["km_total"]) < 0.5
    assert s["amount_total"] == s["km_total"]  # stawka 1 zł/km
    r = data["referees"][0]
    assert r["km"] > 0 and r["trips"] and r["trips"][0]["route"].startswith("dom → ")
    assert sum(m["matches"] for m in data["by_month"]) == 24
    assert sum(lv["matches"] for lv in data["by_level"]) == 24

    x = c.get("/api/statistics/mileage.xlsx", params=PERIOD)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Kilometrówka", "Przejazdy", "Miesiące"]
    rows = list(wb["Kilometrówka"].iter_rows(min_row=3, values_only=True))
    assert rows[-1][0] == "Razem" and len(rows) == len(data["referees"]) + 1
