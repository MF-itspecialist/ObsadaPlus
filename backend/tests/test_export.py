"""Eksport obsady (Excel, PDF) i publikacja wersji."""

import io

from openpyxl import load_workbook

PERIOD = {"date_from": "2026-10-09", "date_to": "2026-10-12"}


def _confirmed(c):
    preview = c.post("/api/optimization/generate", json=PERIOD).json()
    c.post("/api/optimization/confirm", json={"run_id": preview["run_id"]})
    return preview


def test_export_xlsx_lists_every_match_and_marks_missing_roles(seeded_client):
    c = seeded_client
    _confirmed(c)
    r = c.get("/api/assignments/export.xlsx", params=PERIOD)
    assert r.status_code == 200
    assert 'filename="obsada_2026-10-09_2026-10-12.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Obsada"]
    header = [cell.value for cell in ws[2]]
    assert header[:3] == ["Data", "Dzień", "Godzina"]
    assert "Sędzia główny" in header
    data_rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(data_rows) == 24
    main_col = header.index("Sędzia główny")
    assert any(row[main_col] == "— brak —" for row in data_rows), "mecz bez obsady jest oznaczony"
    referees = wb["Sędziowie"]
    assert any(str(row[0]).endswith("— razem") for row in referees.iter_rows(min_row=2, values_only=True))


def test_export_pdf_renders_polish_text(seeded_client):
    c = seeded_client
    _confirmed(c)
    r = c.get("/api/assignments/export.pdf", params=PERIOD)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert b"DejaVu" in r.content  # osadzona czcionka z polskimi znakami


def test_export_rejects_reversed_period(client):
    r = client.get("/api/assignments/export.pdf", params={"date_from": "2026-10-12", "date_to": "2026-10-09"})
    assert r.status_code == 422


def test_publication_versions_and_changes_since_publish(seeded_client):
    c = seeded_client
    assert c.get("/api/publications/latest", params=PERIOD).json() is None
    _confirmed(c)

    first = c.post("/api/publications", json=PERIOD).json()
    assert first["version"] == 1
    assert first["match_count"] == 24
    assert first["changed_matches"] == []

    # ręczna zmiana po publikacji: usunięcie jednego przypisania
    a = c.get("/api/assignments", params=PERIOD).json()[0]
    c.delete(f"/api/assignments/{a['id']}")
    latest = c.get("/api/publications/latest", params=PERIOD).json()
    assert latest["id"] == first["id"]
    assert latest["changed_matches"] == [a["match_id"]]

    # opublikowana kopia się nie zmienia
    full = c.get(f"/api/publications/{first['id']}").json()
    row = next(r for r in full["rows"] if r["match_id"] == a["match_id"])
    assert any(m["referee_id"] == a["referee_id"] for m in row["crew"])

    second = c.post("/api/publications", json=PERIOD).json()
    assert second["version"] == 2
    assert second["changed_matches"] == []
    assert c.get("/api/publications/latest", params=PERIOD).json()["id"] == second["id"]

    pdf = c.get(f"/api/publications/{first['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert c.get("/api/publications/999").status_code == 404
